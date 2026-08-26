from datetime import date

import pytest
from openpyxl import load_workbook

from atlas_eval.models import (
    Bank, Checks, Question, QuestionType, Stance, Verification, VerificationStatus,
)
from atlas_eval.review import (
    ACCURATE_CHOICES, REVIEW_COLUMNS, REVIEW_SHEET, TRACKING_COLUMNS, TRACKING_SHEET,
    export_review,
)


def _q(qid, status=VerificationStatus.UNVERIFIED, **over):
    base = dict(
        id=qid, text=f"question text for {qid}", type=QuestionType.TERM_REDIRECT,
        type_note="term-redirect", provenance="product-brief/employer-charges-phase-2",
        expected_stance=Stance.REDIRECT,
        ground_truth="Maps to the Accelerated Collection Process.",
        checks=Checks(required_all=["ACP"], forbidden=["ETA-227"],
                      expect_citations=["UIMB0730", "L2ACPD"]),
        verification=Verification(status=status),
    )
    base.update(over)
    return Question(**base)


def _bank(*questions, version="v3"):
    return Bank(version=version, agent="a", corpus="c", questions=list(questions))


def test_export_writes_both_sheets_with_exact_headers(tmp_path):
    out = tmp_path / "review.xlsx"
    assert export_review([_bank(_q("v3-Q1"))], out) == 1

    wb = load_workbook(out)
    assert wb.sheetnames == [REVIEW_SHEET, TRACKING_SHEET]
    review = wb[REVIEW_SHEET]
    assert tuple(c.value for c in review[1]) == REVIEW_COLUMNS
    assert tuple(c.value for c in wb[TRACKING_SHEET][1]) == TRACKING_COLUMNS


def test_review_row_carries_question_and_ground_truth(tmp_path):
    out = tmp_path / "review.xlsx"
    export_review([_bank(_q("v3-Q1"))], out)
    row = {k: v for k, v in zip(REVIEW_COLUMNS, [c.value for c in load_workbook(out)[REVIEW_SHEET][2]])}
    assert row["id"] == "v3-Q1"
    assert row["question"] == "question text for v3-Q1"
    assert row["type"] == "term_redirect"
    assert "Accelerated Collection Process" in row["expected answer"]
    assert row["accurate?"] in (None, "")
    assert row["reviewer"] in (None, "")


def test_tracking_sheet_flattens_check_terms(tmp_path):
    out = tmp_path / "review.xlsx"
    export_review([_bank(_q("v3-Q1"))], out)
    row = {k: v for k, v in zip(TRACKING_COLUMNS, [c.value for c in load_workbook(out)[TRACKING_SHEET][2]])}
    assert row["bank"] == "v3"
    assert row["type_note"] == "term-redirect"
    assert row["required_all"] == "ACP"
    assert row["forbidden"] == "ETA-227"
    assert row["expect_citations"] == "UIMB0730; L2ACPD"
    assert row["current status"] == "unverified"
    assert row["expected stance"] == "redirect"


def test_status_filter_selects_a_subset(tmp_path):
    bank = _bank(
        _q("v3-Q1", VerificationStatus.UNVERIFIED),
        _q("v3-Q2", VerificationStatus.VERIFIED,
           verification=Verification(status=VerificationStatus.VERIFIED,
                                     verified_by="Oscar", verified_on=date(2026, 8, 26))),
        _q("v3-Q3", VerificationStatus.NEEDS_SME),
    )
    out = tmp_path / "review.xlsx"
    n = export_review([bank], out, statuses={VerificationStatus.UNVERIFIED,
                                             VerificationStatus.NEEDS_SME})
    assert n == 2
    ids = [r[0].value for r in load_workbook(out)[REVIEW_SHEET].iter_rows(min_row=2)]
    assert ids == ["v3-Q1", "v3-Q3"]


def test_no_filter_exports_everything(tmp_path):
    bank = _bank(_q("v3-Q1"), _q("v3-Q2", VerificationStatus.VERIFIED,
                                 verification=Verification(status=VerificationStatus.VERIFIED,
                                                           verified_by="Oscar")))
    out = tmp_path / "review.xlsx"
    assert export_review([bank], out) == 2


def test_rejected_questions_carry_their_prior_verdict(tmp_path):
    q = _q("v3-Q4", verification=Verification(
        status=VerificationStatus.REJECTED, verified_by="Oscar",
        verified_on=date(2026, 8, 26), notes="No knowable answer; Oscar did not know it either."))
    out = tmp_path / "review.xlsx"
    export_review([_bank(q)], out)
    review = load_workbook(out)[REVIEW_SHEET]
    row = {k: v for k, v in zip(REVIEW_COLUMNS, [c.value for c in review[2]])}
    assert row["accurate?"] == "no"
    assert row["reviewer"] == "Oscar"
    assert "No knowable answer" in row["notes"]


def test_export_of_empty_selection_still_writes_headers(tmp_path):
    out = tmp_path / "review.xlsx"
    assert export_review([_bank(_q("v3-Q1", VerificationStatus.VERIFIED,
                                  verification=Verification(status=VerificationStatus.VERIFIED,
                                                            verified_by="O")))],
                         out, statuses={VerificationStatus.REJECTED}) == 0
    wb = load_workbook(out)
    assert tuple(c.value for c in wb[REVIEW_SHEET][1]) == REVIEW_COLUMNS
    assert wb[REVIEW_SHEET].max_row == 1


def test_accurate_choices_are_the_three_documented_values():
    assert ACCURATE_CHOICES == ("yes", "no", "unsure")
